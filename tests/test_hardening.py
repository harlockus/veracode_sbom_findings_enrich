"""Cross-platform and security behavior that does not call Veracode."""

import gzip
import json
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import load_workbook

from sbom_findings.client import (
    VeracodeClient,
    VeracodeError,
    _gunzip_limited,
    _urllib_sender,
    assert_api_path,
)
from sbom_findings.config import Settings, load_settings
from sbom_findings.export import write_findings_workbook, write_json
from sbom_findings.pipeline import _slug, _write_raw_sboms
from sbom_findings.registry import search_maven_ref


def test_user_agent_stays_on_one_header_line(monkeypatch):
    monkeypatch.setenv("VERACODE_API_KEY_ID", "abcd")
    monkeypatch.setenv("VERACODE_API_KEY_SECRET", "abcd")
    monkeypatch.setenv("VERACODE_USER_AGENT", "SBOM\r\nX-Injected: yes")
    settings = load_settings()
    assert "\r" not in settings.user_agent
    assert "\n" not in settings.user_agent
    assert settings.user_agent == "SBOM X-Injected: yes"

    monkeypatch.setenv("VERACODE_USER_AGENT", "\r\n\t")
    assert load_settings().user_agent == "SBOM-Findings/1.0"


def test_solr_metacharacters_are_not_queried():
    assert search_maven_ref('log4j" OR a:"x', "1.2.8") is None
    assert search_maven_ref("log4j", '1.2.8" OR v:"1') is None
    assert search_maven_ref("log4j", "") is None


def test_slug_is_a_legal_filename_on_windows():
    assert _slug("CON") == "file-CON"
    assert _slug("AUX") == "file-AUX"
    assert _slug("COM1") == "file-COM1"
    hostile = _slug(r"..\..\CON:evil/passwd")
    assert "/" not in hostile
    assert "\\" not in hostile
    assert ":" not in hostile
    assert ".." not in hostile
    assert not hostile.endswith(".")
    assert _slug("   ") == "sbom"


def test_raw_sbom_file_stays_inside_the_output_directory(tmp_path: Path):
    documents = [
        {
            "project_name": r"..\..\CON:evil",
            "scan_mode": "AGENT",
            "format": "cyclonedx",
            "document": {"components": []},
        }
    ]
    _write_raw_sboms(tmp_path, documents)
    written = list((tmp_path / "sbom").glob("*.json"))
    assert len(written) == 1
    assert written[0].parent == tmp_path / "sbom"
    assert "\\" not in documents[0]["file"]
    assert "/" in documents[0]["file"]


def test_excel_formula_text_is_stored_as_text(tmp_path: Path):
    path = tmp_path / "findings.xlsx"
    write_findings_workbook(
        path,
        [
            {
                "component_name": "=cmd|'/c calc'!A1",
                "description": " +quoted",
                "cve": "@SUM(A1)",
                "component_version": "-1.2.3",
                "cvss3_score": 9.8,
            }
        ],
        [{"metric": "findings", "value": 1}],
    )
    sheet = load_workbook(path)["Findings"]
    headers = [cell.value for cell in sheet[1]]
    values = {
        headers[col - 1]: sheet.cell(2, col).value
        for col in range(1, len(headers) + 1)
        if sheet.cell(2, col).value is not None
    }
    assert values["Component"].startswith("'=")
    assert values["Description"].startswith("'")
    assert values["CVE"].startswith("'@")
    assert values["Version"].startswith("'-")
    assert values["CVSS v3"] == 9.8


def test_json_uses_lf_newlines(tmp_path: Path):
    path = tmp_path / "report.json"
    write_json(path, {"ok": True})
    raw = path.read_bytes()
    assert b"\r\n" not in raw
    assert json.loads(raw) == {"ok": True}


def test_encoded_parent_segments_are_refused():
    assert assert_api_path("/appsec/v1/applications") == "/appsec/v1/applications"
    for path in ("/appsec/%2e%2e/admin", "/appsec/%2E%2E/admin", "/appsec/%00", r"/appsec/\windows"):
        try:
            assert_api_path(path)
        except VeracodeError:
            continue
        raise AssertionError(path)


def test_redirect_is_not_followed_and_is_not_success():
    seen: list[str] = []

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            seen.append(self.path)
            if self.path == "/start":
                self.send_response(302)
                self.send_header("Location", "/secret")
                self.end_headers()
                return
            self.send_response(200)
            self.end_headers()
            self.wfile.write(b"secret")

        def log_message(self, format, *args):
            return

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    port = server.server_address[1]
    try:
        status, _headers, _body = _urllib_sender(
            f"http://127.0.0.1:{port}/start",
            {"Authorization": "VERACODE-HMAC-SHA-256 id=test"},
            5,
        )
        assert status == 302
        assert seen == ["/start"]
        settings = Settings("abcd", "abcd", "api.veracode.com", "us", "test", "", "")

        def sender(url, headers, timeout):
            return 302, {"Location": "https://example.invalid/"}, b""

        client = VeracodeClient(settings, sender=sender, sleeper=lambda _delay: None)
        try:
            client.get_json("/appsec/v1/applications")
        except VeracodeError as exc:
            assert exc.status == 302
        else:
            raise AssertionError("redirect was treated as success")
    finally:
        server.shutdown()
        server.server_close()


def test_gzip_expansion_is_capped():
    raw = gzip.compress(b"A" * 4096)
    assert _gunzip_limited(raw, limit=4096) == b"A" * 4096
    try:
        _gunzip_limited(raw, limit=100)
    except VeracodeError as exc:
        assert "size" in str(exc)
    else:
        raise AssertionError("gzip expansion was accepted")


def test_non_https_remote_url_is_refused():
    try:
        _urllib_sender("http://example.com/appsec", {}, 5)
    except VeracodeError as exc:
        assert "HTTPS" in str(exc)
    else:
        raise AssertionError("plain HTTP was accepted")
