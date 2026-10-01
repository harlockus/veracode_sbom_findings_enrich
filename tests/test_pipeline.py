import json
from pathlib import Path

from openpyxl import load_workbook

from sbom_findings.client import VeracodeError
from sbom_findings.config import Settings
from sbom_findings.pipeline import CollectRequest, collect
from sbom_findings.client import VeracodeClient

APP = "11111111-1111-1111-1111-111111111111"
PROJECT = "22222222-2222-2222-2222-222222222222"
WORKSPACE = "33333333-3333-3333-3333-333333333333"
ISSUE = "44444444-4444-4444-4444-444444444444"


def _page(key, rows):
    return {
        "_embedded": {key: rows},
        "page": {"number": 0, "size": 100, "total_elements": len(rows), "total_pages": 1},
    }


UPLOAD_SBOM = {
    "bomFormat": "CycloneDX",
    "specVersion": "1.4",
    "components": [
        {
            "bom-ref": "pkg:npm/console-io@2.6.3",
            "type": "library",
            "name": "console-io",
            "version": "2.6.3",
            "purl": "pkg:npm/console-io@2.6.3",
            "licenses": [{"license": {"id": "MIT"}}],
        }
    ],
    "vulnerabilities": [
        {
            "id": "CVE-2016-10532",
            "description": "already in findings",
            "ratings": [{"score": 9.8, "severity": "high"}],
            "affects": [{"ref": "pkg:npm/console-io@2.6.3"}],
        },
        {
            "id": "CVE-SBOM-ONLY",
            "description": "only in the bill of materials",
            "ratings": [{"score": 5.0, "severity": "medium"}],
            "affects": [{"ref": "pkg:npm/console-io@2.6.3"}],
        },
    ],
    "dependencies": [{"ref": "app", "dependsOn": ["pkg:npm/console-io@2.6.3"]}],
}


class Scripted(VeracodeClient):
    def __init__(self, *, linked_status: int = 200):
        settings = Settings("aa", "bb", "api.veracode.com", "us", "test", "", "")
        super().__init__(settings, sender=lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("network")))
        self.linked_status = linked_status
        self.calls = []

    def get_json(self, path, query=None):
        base = path.split("?", 1)[0]
        query = dict(query or {})
        self.calls.append((base, query))
        if base == "/appsec/v1/applications":
            return _page(
                "applications",
                [{"guid": APP, "profile": {"name": "verademo"}}],
            )
        if base.endswith("/findings"):
            return _page(
                "findings",
                [
                    {
                        "scan_type": "SCA",
                        "description": "console-io auth bypass",
                        "issue_id": 42,
                        "finding_status": {"status": "OPEN"},
                        "finding_details": {
                            "severity": 5,
                            "component_filename": "console-io",
                            "version": "2.6.3",
                            "language": "JAVASCRIPT",
                            "metadata": {"sca_scan_mode": "UPLOAD", "sca_dep_mode": "DIRECT"},
                            "cve": {"name": "CVE-2016-10532", "cvss3": {"score": 9.8}},
                            "cwe": {"id": 287, "name": "Improper Authentication"},
                        },
                    },
                    {
                        "scan_type": "SCA",
                        "description": "closed console-io finding",
                        "issue_id": 43,
                        "finding_status": {"status": "CLOSED", "resolution": "FIXED"},
                        "finding_details": {
                            "severity": 3,
                            "component_filename": "console-io",
                            "version": "2.6.3",
                            "language": "JAVASCRIPT",
                            "metadata": {"sca_scan_mode": "UPLOAD", "sca_dep_mode": "DIRECT"},
                            "cve": {"name": "CVE-CLOSED-ONLY", "cvss3": {"score": 4.0}},
                        },
                    },
                    {
                        "scan_type": "SCA",
                        "description": "mitigated but still open",
                        "issue_id": 44,
                        "finding_status": {"status": "OPEN", "resolution": "MITIGATED"},
                        "finding_details": {
                            "severity": 4,
                            "component_filename": "left-pad",
                            "version": "1.0.0",
                            "language": "JAVASCRIPT",
                            "metadata": {"sca_scan_mode": "UPLOAD", "sca_dep_mode": "DIRECT"},
                            "cve": {"name": "CVE-MITIGATED-OPEN", "cvss3": {"score": 6.1}},
                        },
                    },
                ],
            )
        if base.endswith("/cyclonedx") and query.get("type") == "application" and query.get("linked") == "true":
            return {"bomFormat": "CycloneDX", "components": [], "vulnerabilities": [], "dependencies": []}
        if base.endswith("/cyclonedx") and query.get("type") == "application":
            return UPLOAD_SBOM
        if base.endswith("/cyclonedx") and query.get("type") == "agent":
            return {
                "bomFormat": "CycloneDX",
                "components": [
                    {
                        "bom-ref": "pkg:maven/net.minidev/json-smart@1.3.1",
                        "name": "json-smart",
                        "version": "1.3.1",
                        "purl": "pkg:maven/net.minidev/json-smart@1.3.1",
                    }
                ],
                "vulnerabilities": [],
                "dependencies": [],
            }
        if base.endswith("/projects") and "/applications/" in base:
            if self.linked_status != 200:
                raise VeracodeError("denied", status=self.linked_status, path=base)
            return {
                "linked_projects": [
                    {
                        "id": PROJECT,
                        "name": "verademo-src",
                        "workspace": {"id": WORKSPACE, "name": "demo-workspace"},
                    }
                ]
            }
        if base.endswith("/issues"):
            return _page(
                "issues",
                [
                    {
                        "id": ISSUE,
                        "issue_status": "open",
                        "issue_type": "vulnerability",
                        "project_id": PROJECT,
                        "project_name": "verademo-src",
                        "library": {
                            "id": "maven:net.minidev:json-smart:1.3.1:",
                            "name": "json-smart",
                            "version": "1.3.1",
                            "direct": False,
                            "transitive": True,
                        },
                        "vulnerability": {"cve": "CVE-2023-1370", "cvss3_score": 7.5, "title": "overflow"},
                    },
                    {
                        "id": "55555555-5555-5555-5555-555555555555",
                        "issue_status": "fixed",
                        "issue_type": "vulnerability",
                        "project_id": PROJECT,
                        "project_name": "verademo-src",
                        "library": {
                            "id": "npm:old-lib::1.0.0:",
                            "name": "old-lib",
                            "version": "1.0.0",
                        },
                        "vulnerability": {"cve": "CVE-FIXED-ONLY", "cvss3_score": 3.1, "title": "fixed"},
                    },
                ],
            )
        if base == f"/srcclr/v3/issues/{ISSUE}":
            return {
                "id": ISSUE,
                "fix_info": {"fixed_version": "2.4.2", "latest_safe_version": "2.5.1"},
                "library": {"id": "maven:net.minidev:json-smart:1.3.1:", "latest_version": "2.5.1"},
            }
        if base.startswith("/srcclr/v3/component-activity/"):
            if "json-smart" in path:
                return {
                    "safe_versions": ["2.4.2", "2.5.1"],
                    "library": {"latest_release_version": "2.5.1"},
                }
            return {"safe_versions": ["3.0.0"], "library": {"latest_release_version": "3.1.0"}}
        raise AssertionError(f"unexpected call {base} {query}")


def test_collect_writes_json_and_xlsx_with_safe_versions(tmp_path: Path):
    client = Scripted()
    result = collect(
        client,
        CollectRequest(app_names=["verademo"], out_dir=tmp_path, enrich=True),
    )
    findings = json.loads(result.findings_json.read_text())["findings"]
    by_cve = {row["cve"]: row for row in findings}
    assert by_cve["CVE-2016-10532"]["recommended_safe_version"] == "3.0.0"
    assert by_cve["CVE-2016-10532"]["library_ref"] == "npm:console-io::2.6.3:"
    assert "sbom" in by_cve["CVE-2016-10532"]["sources"]
    assert by_cve["CVE-2023-1370"]["fixed_version"] == "2.4.2"
    assert by_cve["CVE-2023-1370"]["latest_safe_version"] == "2.5.1"
    assert by_cve["CVE-2023-1370"]["recommended_safe_version"] == "2.5.1"
    assert by_cve["CVE-2023-1370"]["safe_versions"] == ["2.4.2", "2.5.1"]
    assert by_cve["CVE-2023-1370"]["dependency_mode"] == "TRANSITIVE"
    assert by_cve["CVE-SBOM-ONLY"]["source"] == "sbom"
    assert by_cve["CVE-MITIGATED-OPEN"]["issue_status"] == "OPEN"
    assert "CVE-CLOSED-ONLY" not in by_cve
    assert "CVE-FIXED-ONLY" not in by_cve
    issue_calls = [query for base, query in client.calls if base.endswith("/issues")]
    assert issue_calls
    assert all(query.get("status") == "open" for query in issue_calls)

    sbom = json.loads(result.sbom_json.read_text())
    assert any(row["name"] == "console-io" for row in sbom["components"])
    assert any(doc["file"].endswith(".json") for doc in sbom["documents"])
    assert "\\" not in sbom["documents"][0]["file"]
    assert (tmp_path / sbom["documents"][0]["file"]).is_file()

    findings_book = load_workbook(result.findings_xlsx)
    assert findings_book.sheetnames == ["Findings", "Safe versions", "Enrichment gaps", "Summary"]
    header = [cell.value for cell in next(findings_book["Findings"].iter_rows(max_row=1))]
    assert "Recommended safe version" in header
    assert "CVE" in header

    sbom_book = load_workbook(result.sbom_xlsx)
    assert sbom_book.sheetnames == ["Components", "Vulnerabilities", "Dependencies", "Licenses", "Summary"]
    assert result.manifest_json.is_file()


def test_forbidden_agent_api_still_writes_upload_results(tmp_path: Path):
    client = Scripted(linked_status=403)
    result = collect(client, CollectRequest(app_names=["verademo"], out_dir=tmp_path))
    findings = json.loads(result.findings_json.read_text())["findings"]
    assert any(row["cve"] == "CVE-2016-10532" for row in findings)
    assert any("access denied" in warning for warning in result.warnings)
