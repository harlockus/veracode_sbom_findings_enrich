"""Write the SBOM and findings results as JSON and Excel workbooks."""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_ILLEGAL = re.compile(r"[\000-\010\013\014\016-\037]")

FINDING_COLUMNS = [
    ("sources", "Sources"),
    ("scan_mode", "Scan mode"),
    ("dependency_mode", "Dependency"),
    ("application_name", "Application"),
    ("application_guid", "Application GUID"),
    ("workspace_name", "Workspace"),
    ("project_name", "Project"),
    ("project_branch", "Branch"),
    ("issue_id", "Issue id"),
    ("issue_status", "Status"),
    ("resolution", "Resolution"),
    ("first_found_date", "First found"),
    ("last_seen_date", "Last seen"),
    ("violates_policy", "Violates policy"),
    ("policy_severity", "Severity"),
    ("cve", "CVE"),
    ("cwe_id", "CWE"),
    ("cwe_name", "CWE name"),
    ("cvss2_score", "CVSS v2"),
    ("cvss3_score", "CVSS v3"),
    ("cvss3_severity", "CVSS v3 severity"),
    ("cvss3_vector", "CVSS v3 vector"),
    ("epss_score", "EPSS"),
    ("epss_percentile", "EPSS percentile"),
    ("exploit_observed", "Exploit observed"),
    ("exploit_source", "Exploit source"),
    ("component_name", "Component"),
    ("component_version", "Version"),
    ("language", "Language"),
    ("licenses", "Licenses"),
    ("component_paths", "Paths"),
    ("vulnerable_method", "Vulnerable method"),
    ("vulnerable_methods", "Vulnerable methods"),
    ("library_ref", "Library reference"),
    ("fixed_version", "Fixed version"),
    ("latest_safe_version", "Latest safe version"),
    ("safe_versions", "Safe versions"),
    ("recommended_safe_version", "Recommended safe version"),
    ("latest_release_version", "Latest release"),
    ("enrichment_status", "Enrichment"),
    ("enrichment_detail", "Enrichment detail"),
    ("description", "Description"),
]

COMPONENT_COLUMNS = [
    ("scan_mode", "Scan mode"),
    ("sbom_format", "Format"),
    ("application_name", "Application"),
    ("workspace_name", "Workspace"),
    ("project_name", "Project"),
    ("name", "Component"),
    ("group", "Group"),
    ("version", "Version"),
    ("purl", "Package URL"),
    ("type", "Type"),
    ("licenses", "Licenses"),
    ("library_ref", "Library reference"),
    ("library_ref_source", "Reference source"),
    ("bom_ref", "BOM ref"),
    ("hashes", "Hashes"),
]

VULN_COLUMNS = [
    ("scan_mode", "Scan mode"),
    ("sbom_format", "Format"),
    ("application_name", "Application"),
    ("project_name", "Project"),
    ("vulnerability_id", "Vulnerability"),
    ("severity", "Severity"),
    ("score", "Score"),
    ("vector", "Vector"),
    ("cwes", "CWEs"),
    ("affected_components", "Affected components"),
    ("source_name", "Source"),
    ("source_url", "URL"),
    ("description", "Description"),
]

DEP_COLUMNS = [
    ("scan_mode", "Scan mode"),
    ("sbom_format", "Format"),
    ("application_name", "Application"),
    ("project_name", "Project"),
    ("ref", "Component"),
    ("depends_on", "Depends on"),
]


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(payload, indent=2, ensure_ascii=False, default=_json_default) + "\n"
    path.write_text(text, encoding="utf-8", newline="\n")


def write_findings_workbook(path: Path, findings: list[dict[str, Any]], summary: list[dict[str, Any]]) -> None:
    book = Workbook()
    _write_sheet(book.active, "Findings", FINDING_COLUMNS, findings)
    safe_rows = [row for row in findings if row.get("cve") or row.get("issue_type") == "vulnerability"]
    _write_sheet(book.create_sheet("Safe versions"), "Safe versions", FINDING_COLUMNS, safe_rows)
    gaps = [
        row
        for row in findings
        if row.get("enrichment_status") not in {"enriched", ""}
    ]
    _write_sheet(
        book.create_sheet("Enrichment gaps"),
        "Enrichment gaps",
        FINDING_COLUMNS,
        gaps,
    )
    _write_sheet(
        book.create_sheet("Summary"),
        "Summary",
        [("metric", "Metric"), ("value", "Value")],
        summary,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def write_sbom_workbook(
    path: Path,
    *,
    components: list[dict[str, Any]],
    vulnerabilities: list[dict[str, Any]],
    dependencies: list[dict[str, Any]],
    summary: list[dict[str, Any]],
) -> None:
    book = Workbook()
    _write_sheet(book.active, "Components", COMPONENT_COLUMNS, components)
    _write_sheet(book.create_sheet("Vulnerabilities"), "Vulnerabilities", VULN_COLUMNS, vulnerabilities)
    _write_sheet(book.create_sheet("Dependencies"), "Dependencies", DEP_COLUMNS, dependencies)
    licenses = _license_rows(components)
    _write_sheet(
        book.create_sheet("Licenses"),
        "Licenses",
        [
            ("license", "License"),
            ("component", "Component"),
            ("version", "Version"),
            ("application_name", "Application"),
            ("project_name", "Project"),
            ("scan_mode", "Scan mode"),
        ],
        licenses,
    )
    _write_sheet(
        book.create_sheet("Summary"),
        "Summary",
        [("metric", "Metric"), ("value", "Value")],
        summary,
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def _license_rows(components: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    for component in components:
        licenses = component.get("licenses") or []
        if isinstance(licenses, str):
            licenses = [licenses]
        for license_name in licenses:
            if not license_name:
                continue
            rows.append(
                {
                    "license": license_name,
                    "component": component.get("name"),
                    "version": component.get("version"),
                    "application_name": component.get("application_name"),
                    "project_name": component.get("project_name"),
                    "scan_mode": component.get("scan_mode"),
                }
            )
    return rows


def _write_sheet(sheet, title: str, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    sheet.title = title[:31]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    wrap = Alignment(wrap_text=True, vertical="top")
    for col, (_, label) in enumerate(columns, start=1):
        cell = sheet.cell(1, col, label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    for row_index, row in enumerate(rows, start=2):
        for col, (key, _) in enumerate(columns, start=1):
            cell = sheet.cell(row_index, col, _cell(row.get(key)))
            cell.alignment = wrap
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{max(len(rows) + 1, 1)}"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col, (key, label) in enumerate(columns, start=1):
        width = max(len(label), 12)
        for row in rows[:200]:
            rendered = _cell(row.get(key))
            text = "" if rendered is None else str(rendered)
            width = max(width, min(len(text), 60))
        sheet.column_dimensions[get_column_letter(col)].width = min(width + 2, 62)
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _cell(value: Any) -> str | int | float | bool | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    if isinstance(value, (list, tuple)):
        text = "; ".join(str(item) for item in value)
    else:
        text = str(value)
    text = _ILLEGAL.sub("", text).replace("\r", "")
    # A leading formula character is attacker-controlled text from a component
    # name or description. The apostrophe stores it as text in Excel.
    if text.lstrip(" \t").startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _json_default(value: Any) -> str:
    return str(value)
